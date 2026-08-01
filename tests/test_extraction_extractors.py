"""Unit tests for the per-format extractors in app/core/extraction/.

Test fixtures are generated on the fly (via PyMuPDF/python-docx/plain
writes) rather than committed as binary files -- fully synthetic,
reproducible, and reviewable as ordinary Python code.
"""

from __future__ import annotations

from email.message import EmailMessage
from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument

from app.core.extraction import docx as docx_extractor
from app.core.extraction import email as email_extractor
from app.core.extraction import pdf as pdf_extractor
from app.core.extraction import text as text_extractor
from app.core.extraction.dispatcher import get_extractor, is_image_extension
from app.core.extraction.needs_ocr import page_needs_ocr


# --- needs_ocr heuristic -----------------------------------------------


def test_page_needs_ocr_false_when_no_visual_content():
    assert page_needs_ocr(None, has_visual_content=False) is False
    assert page_needs_ocr("", has_visual_content=False) is False


def test_page_needs_ocr_true_for_image_with_little_text():
    assert page_needs_ocr("a b", has_visual_content=True) is True
    assert page_needs_ocr(None, has_visual_content=True) is True


def test_page_needs_ocr_false_for_image_with_substantial_text():
    substantial = " ".join(f"word{i}" for i in range(20))
    assert page_needs_ocr(substantial, has_visual_content=True) is False


# --- dispatcher ----------------------------------------------------------


@pytest.mark.parametrize(
    "filename", ["a.pdf", "A.PDF", "b.docx", "c.txt", "d.rtf", "e.eml"]
)
def test_dispatcher_recognizes_supported_extensions(filename: str):
    assert get_extractor(filename) is not None


@pytest.mark.parametrize("filename", ["a.csv", "b.xlsx", "c.doc", "d.msg", "no-extension"])
def test_dispatcher_returns_none_for_unsupported_extensions(filename: str):
    assert get_extractor(filename) is None


@pytest.mark.parametrize("filename", ["a.jpg", "A.JPEG", "b.png", "c.tif", "d.tiff"])
def test_dispatcher_recognizes_image_extensions(filename: str):
    assert is_image_extension(filename) is True
    assert get_extractor(filename) is None  # images have no text extractor


# --- PDF -------------------------------------------------------------------


def _make_pdf_with_text(path: Path, text: str) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    doc.save(str(path))
    doc.close()


def _make_pdf_image_only(path: Path) -> None:
    """A page with an embedded image and no text -- should need OCR."""
    doc = fitz.open()
    page = doc.new_page()
    pixmap = fitz.Pixmap(fitz.csRGB, (0, 0, 20, 20))
    pixmap.set_rect(pixmap.irect, (200, 0, 0))
    page.insert_image(fitz.Rect(72, 72, 200, 200), pixmap=pixmap)
    doc.save(str(path))
    doc.close()


def test_pdf_extract_returns_native_text(tmp_path: Path):
    pdf_path = tmp_path / "sample.pdf"
    _make_pdf_with_text(pdf_path, "Hello World from a real PDF page")

    result = pdf_extractor.extract(pdf_path)

    assert len(result.pages) == 1
    page = result.pages[0]
    assert page.page_number == 1
    assert "Hello World" in (page.text or "")
    assert page.extraction_method == "native"
    assert page.char_count > 0
    assert page.needs_ocr is False


def test_pdf_extract_flags_image_only_page_as_needing_ocr(tmp_path: Path):
    pdf_path = tmp_path / "scanned.pdf"
    _make_pdf_image_only(pdf_path)

    result = pdf_extractor.extract(pdf_path)

    assert len(result.pages) == 1
    assert result.pages[0].needs_ocr is True
    assert result.pages[0].char_count == 0


def test_pdf_extract_handles_multiple_pages(tmp_path: Path):
    pdf_path = tmp_path / "multi.pdf"
    doc = fitz.open()
    for i in range(3):
        page = doc.new_page()
        page.insert_text((72, 72), f"Page number {i + 1}")
    doc.save(str(pdf_path))
    doc.close()

    result = pdf_extractor.extract(pdf_path)

    assert len(result.pages) == 3
    assert [p.page_number for p in result.pages] == [1, 2, 3]
    assert "Page number 2" in (result.pages[1].text or "")


def test_pdf_extract_raises_on_corrupt_file(tmp_path: Path):
    corrupt = tmp_path / "corrupt.pdf"
    corrupt.write_bytes(b"this is not a valid PDF file at all")

    with pytest.raises(Exception):
        pdf_extractor.extract(corrupt)


# --- DOCX --------------------------------------------------------------


def test_docx_extract_returns_paragraph_text(tmp_path: Path):
    docx_path = tmp_path / "sample.docx"
    doc = DocxDocument()
    doc.add_paragraph("First paragraph of the document.")
    doc.add_paragraph("Second paragraph, with more content.")
    doc.save(str(docx_path))

    result = docx_extractor.extract(docx_path)

    assert len(result.pages) == 1
    page = result.pages[0]
    assert page.page_number == 1
    assert "First paragraph" in page.text
    assert "Second paragraph" in page.text
    assert page.extraction_method == "native"
    assert page.needs_ocr is False


def test_docx_extract_raises_on_corrupt_file(tmp_path: Path):
    corrupt = tmp_path / "corrupt.docx"
    corrupt.write_bytes(b"not a real docx zip archive")

    with pytest.raises(Exception):
        docx_extractor.extract(corrupt)


# --- plain text / RTF ----------------------------------------------------


def test_text_extract_plain_text(tmp_path: Path):
    txt_path = tmp_path / "sample.txt"
    txt_path.write_text("Plain text content here.\nSecond line.")

    result = text_extractor.extract_plain_text(txt_path)

    assert len(result.pages) == 1
    assert "Plain text content" in result.pages[0].text
    assert result.pages[0].needs_ocr is False


def test_text_extract_rtf(tmp_path: Path):
    rtf_path = tmp_path / "sample.rtf"
    rtf_path.write_text(r"{\rtf1\ansi Hello from RTF}")

    result = text_extractor.extract_rtf(rtf_path)

    assert len(result.pages) == 1
    assert "Hello from RTF" in result.pages[0].text


# --- Email -----------------------------------------------------------------


def _write_eml(path: Path, *, with_attachment: bool = False) -> None:
    message = EmailMessage()
    message["From"] = "sender@example.com"
    message["To"] = "recipient@example.com"
    message["Subject"] = "Test Subject Line"
    message["Date"] = "Mon, 1 Jan 2024 12:00:00 -0000"
    message.set_content("This is the body of the test email.")
    if with_attachment:
        message.add_attachment(
            b"attachment file content",
            maintype="text",
            subtype="plain",
            filename="attached.txt",
        )
    path.write_bytes(bytes(message))


def test_email_extract_headers_and_body(tmp_path: Path):
    eml_path = tmp_path / "sample.eml"
    _write_eml(eml_path)

    result = email_extractor.extract(eml_path)

    assert len(result.pages) == 1
    text = result.pages[0].text
    assert "sender@example.com" in text
    assert "Test Subject Line" in text
    assert "This is the body" in text
    assert result.attachments == []


def test_email_extract_returns_attachments_without_ingesting_them(tmp_path: Path):
    eml_path = tmp_path / "with_attachment.eml"
    _write_eml(eml_path, with_attachment=True)

    result = email_extractor.extract(eml_path)

    assert len(result.attachments) == 1
    attachment = result.attachments[0]
    assert attachment.filename == "attached.txt"
    assert attachment.content == b"attachment file content"
